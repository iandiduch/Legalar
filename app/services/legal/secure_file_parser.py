"""Servicio de validación de seguridad de archivos (MIME sniffing por magic bytes),
protección contra inyecciones de prompts y extracción adaptativa de texto/OCR multimodal.
"""

import base64
import io
import logging
import zipfile
from typing import Any

from fastapi import HTTPException, status
from langchain_core.messages import HumanMessage
import pdfplumber
import docx

from app.core.config import Settings
from app.core.prompt_injection_scanner import scan_for_injection_patterns
from app.services.llm_factory import build_vision_model

logger = logging.getLogger(__name__)


# Firmas binarias estándar (Magic Bytes)
MAGIC_SIGNATURES: dict[str, bytes] = {
    "pdf": b"%PDF-",
    "png": b"\x89PNG\r\n\x1a\n",
    "jpeg": b"\xff\xd8\xff",
    "docx_zip": b"PK\x03\x04",
    "riff": b"RIFF",
}


def validate_file_magic_bytes(
    file_bytes: bytes,
    filename: str,
    max_size: int = 10 * 1024 * 1024,
) -> str:
    """Valida tamaño y firma binaria (magic bytes) del archivo para evitar ejecutables camuflados.
    Retorna el MIME type canónico verificado.
    """
    if len(file_bytes) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El archivo proporcionado está vacío.",
        )

    if len(file_bytes) > max_size:
        raise HTTPException(
            status_code=status.HTTP_413_CONTENT_TOO_LARGE,
            detail=f"El archivo excede el tamaño máximo permitido de {max_size // (1024 * 1024)} MB.",
        )

    # 1. PDF
    if file_bytes.startswith(MAGIC_SIGNATURES["pdf"]):
        return "application/pdf"

    # 2. PNG
    if file_bytes.startswith(MAGIC_SIGNATURES["png"]):
        return "image/png"

    # 3. JPEG / JPG
    if file_bytes.startswith(MAGIC_SIGNATURES["jpeg"]):
        return "image/jpeg"

    # 4. WEBP (RIFF....WEBP)
    if file_bytes.startswith(MAGIC_SIGNATURES["riff"]) and len(file_bytes) >= 12:
        if file_bytes[8:12] == b"WEBP":
            return "image/webp"

    # 5. DOCX (Archivo ZIP con estructura Office Open XML)
    if file_bytes.startswith(MAGIC_SIGNATURES["docx_zip"]):
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                names = z.namelist()
                if any("word/document.xml" in n for n in names):
                    return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except Exception:
            pass

    # 6. Texto plano (TXT / MD)
    # Debe ser decodificable en UTF-8 o Latin-1 y no contener bytes nulos típicos de binarios
    if b"\x00" not in file_bytes[:4096]:
        try:
            file_bytes.decode("utf-8")
            return "text/plain"
        except UnicodeDecodeError:
            try:
                file_bytes.decode("latin-1")
                return "text/plain"
            except Exception:
                pass

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail=(
            f"El archivo '{filename}' fue rechazado por seguridad: su contenido binario "
            "no coincide con ningún formato legal permitido (PDF, DOCX, TXT, PNG, JPG, WEBP)."
        ),
    )


def audit_prompt_injection(text: str) -> None:
    """Escanea el texto extraído buscando patrones de prompt injection o manipulación."""
    matches = scan_for_injection_patterns(text)
    if matches:
        matched_str = ", ".join(matches)
        logger.warning("Intento de Prompt Injection detectado en archivo subido: %s", matched_str)
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "El documento fue rechazado por los sistemas de seguridad: contiene instrucciones "
                f"sospechosas que intentan vulnerar o manipular las directivas del asistente ({matched_str})."
            ),
        )


async def extract_document_text(
    file_bytes: bytes,
    mime_type: str,
    settings: Settings,
    vision_llm: Any = None,
) -> str:
    """Extrae texto de forma adaptativa: local para TXT/DOCX/PDF digital,
    o mediante modelo de visión multimodal para imágenes o PDFs escaneados.
    """
    # Caso 1: Texto plano
    if mime_type.startswith("text/"):
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="replace")

    # Caso 2: DOCX
    if "wordprocessingml" in mime_type:
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())
            for t in doc.tables:
                for row in t.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        parts.append(row_text)
            return "\n\n".join(parts)
        except Exception as exc:
            logger.error("Error extrayendo texto de DOCX: %s", exc)
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"No se pudo extraer el texto del archivo DOCX: {exc}",
            ) from exc

    # Caso 3: PDF (Digital o Escaneado)
    if mime_type == "application/pdf":
        extracted_text = ""
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                page_texts = []
                for page in pdf.pages:
                    pt = page.extract_text()
                    if pt:
                        page_texts.append(pt.strip())
                extracted_text = "\n\n".join(page_texts)
        except Exception as exc:
            logger.warning("pdfplumber falló al leer PDF: %s", exc)

        # Si el texto extraído es suficiente, es un PDF digital -> devolvemos directo
        if len(extracted_text.strip()) >= 80:
            return extracted_text.strip()

        # Si tiene menos de 80 caracteres, es un PDF escaneado (fotocopia/imagen)
        logger.info("PDF escaneado detectado (sin capa de texto). Procediendo con visión OCR.")
        # Extraemos la imagen de la primera página con pdfplumber si es posible
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                if pdf.pages:
                    first_page_img = pdf.pages[0].to_image(resolution=150).original
                    buffered = io.BytesIO()
                    first_page_img.save(buffered, format="JPEG")
                    img_bytes = buffered.getvalue()
                    return await _transcribe_image_with_vision(img_bytes, "image/jpeg", settings, vision_llm)
        except Exception as exc:
            logger.warning("Fallo en rasterizado de PDF escaneado: %s", exc)

        if extracted_text.strip():
            return extracted_text.strip()

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El PDF escaneado no contiene texto legible y no pudo ser rasterizado para OCR.",
        )

    # Caso 4: Imágenes directas (PNG, JPEG, WEBP)
    if mime_type.startswith("image/"):
        return await _transcribe_image_with_vision(file_bytes, mime_type, settings, vision_llm)

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail=f"Tipo MIME '{mime_type}' no soportado para extracción de texto.",
    )


async def _transcribe_image_with_vision(
    image_bytes: bytes,
    mime_type: str,
    settings: Settings,
    vision_llm: Any = None,
) -> str:
    """Transcribe fielmente una imagen de contrato o carta documento utilizando el modelo de visión."""
    llm = vision_llm or build_vision_model(settings)
    b64_image = base64.b64encode(image_bytes).decode("utf-8")

    prompt_text = (
        "Eres un transcriptor experto en documentos legales argentinos. "
        "Transcribe con absoluta fidelidad todo el texto visible en este documento o contrato, "
        "conservando la numeración de cláusulas, encabezados, títulos y firmas. "
        "Devuelve únicamente el texto transcrito sin comentarios adicionales."
    )

    message = HumanMessage(
        content=[
            {"type": "text", "text": prompt_text},
            {
                "type": "image_url",
                "image_url": {
                    "url": f"data:{mime_type};base64,{b64_image}",
                },
            },
        ]
    )

    try:
        response = await llm.ainvoke([message])
        transcription = str(response.content).strip()
        if not transcription:
            raise ValueError("El modelo de visión devolvió una transcripción vacía.")
        return transcription
    except Exception as exc:
        logger.error("Fallo al transcribir imagen con modelo de visión (%s): %s", settings.effective_vision_model, exc)
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"No se pudo procesar la imagen con el modelo de visión ({settings.effective_vision_model}): {exc}",
        ) from exc
